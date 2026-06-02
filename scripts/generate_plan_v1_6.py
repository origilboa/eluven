#!/usr/bin/env python3
"""Generate eluven_process_and_plan_v1_6.docx from current project state."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

OUTPUT = Path(__file__).resolve().parent.parent / "docs/spec/eluven_process_and_plan_v1_6.docx"


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value


def build() -> Document:
    doc = Document()

    add_title(doc, "Eluven")
    add_title(doc, "Academic Research Assistant Platform")
    add_title(doc, "Process & Development Plan")
    add_para(doc, "Version 1.6  |  June 2026")

    add_heading(doc, "Changelog — v1.6 (June 2026)", level=2)
    add_bullets(
        doc,
        [
            "MVP completeness audit implemented — critical blockers from v1.5 audit addressed in production",
            "DocumentWorker and WorkflowWorker deployed as ECS Fargate services (infra/workers.ts)",
            "KB upload → indexing verified in production (scripts/validate_mvp_kb_rag.py); plain-text path validated",
            "Migration 006: workflow template seeds (EPR + SPR) and platform instruction seeds (initial_read, final_recommendation)",
            "SPR Assignment/Submission: cluster_type assignment, submission API, Assignments UI",
            "Workflow UI on task detail: start template, poll status, respond to interventions",
            "Word export: GET /tasks/{id}/export/word + frontend download button",
            "Test suites: backend pytest (unit + integration smoke), Jest, Playwright scaffold; CI workflows updated",
            "Logout, Hebrew sidebar labels, docs/infrastructure/master-reference.md updated, ADR 001 for workers",
            "MVP still not fully complete — see Section 1.5 and docs/open-questions.md for remaining validation",
        ],
    )

    add_para(
        doc,
        "Purpose: Master reference for building Eluven. Load at session start with concept doc and data model.",
    )
    add_para(
        doc,
        "Concept doc: academic_research_platform_concept_v1_2.docx. "
        "Data model: eluven_data_model_v1_0.docx. "
        "Prior plan: eluven_process_and_plan_v1_5.docx.",
    )

    add_heading(doc, "1. Project Overview", level=1)
    add_heading(doc, "1.1 What we are building", level=2)
    add_para(
        doc,
        "Eluven is an AI-powered workspace for academic researchers — structured, guided assistance for "
        "reviewing papers, evaluating student submissions, and preparing work for publication.",
    )

    add_heading(doc, "1.2 Platform identity", level=2)
    add_table(
        doc,
        ["Item", "Detail"],
        [
            ["Production app URL", "https://app.eluven.ai"],
            ["Dev credentials", "dev@eluven.ai / devpassword123"],
            ["DNS", "Route 53 zone Z06614832D7DMEV9IQTMN"],
        ],
    )

    add_heading(doc, "1.3 Status at a glance", level=2)
    add_table(
        doc,
        ["Area", "Status"],
        [
            ["Infrastructure", "Production deployed — API + DocumentWorker + WorkflowWorker on ECS"],
            ["Build status", "~75% MVP scope; core paths production-verified"],
            ["Current phase", "Internal validation — full EPR/SPR task runs, instruction quality, RTL"],
            ["External launch", "Not ready — pre-launch decisions (Section 2.2) still open"],
        ],
    )

    add_heading(doc, "1.4 Completed to date (June 2026 — post audit)", level=2)
    add_bullets(
        doc,
        [
            "ECS services: Api, DocumentWorker, WorkflowWorker (SST infra/workers.ts)",
            "KB indexing pipeline verified for .txt uploads; Titan embeddings via Bedrock",
            "Assignments UI (/clusters), submission API, workflow panel, Word export",
            "Alembic migrations 001–006 in production",
            "Automated validation: scripts/validate_mvp_kb_rag.py, scripts/run_mvp_validation.py",
            "Verified: login, dashboard, tasks, threads, AI SSE, KB upload/index",
        ],
    )

    add_heading(doc, "1.5 Remaining gaps (June 2026)", level=2)
    add_bullets(
        doc,
        [
            "One full EPR + one full SPR task not yet validated by internal team (MVP success criteria §8.3)",
            "PDF/DOCX indexing — Unstructured ML deps not in Docker image; .txt validated only",
            "Master instructions drafted for 2 of 11 EPR/SPR thread types only",
            "Rate limiting (10 AI/min) and workflow concurrency cap (3/user) not implemented",
            "Thread document upload UI missing; Instruction Studio cluster/task/thread levels partial",
            "Playwright E2E not in CI; Hebrew/RTL not fully validated",
            "See docs/open-questions.md for full list",
        ],
    )

    add_heading(doc, "1.6 Next sessions in order", level=2)
    add_bullets(
        doc,
        [
            "Internal validation — complete one EPR and one SPR task; document UX friction",
            "Instruction drafting — remaining EPR/SPR thread types; quality bar testing",
            "PDF indexing decision — add Unstructured ML to image vs lighter PDF path",
            "Hebrew/RTL production pass",
            "Wire Playwright into CI; expand E2E to 5 critical paths",
            "Pre-launch decision sessions before external users",
        ],
    )

    add_heading(doc, "3. Tech Stack & Infrastructure", level=1)
    add_table(
        doc,
        ["Component", "Production detail"],
        [
            ["Frontend", "https://app.eluven.ai — CloudFront + Lambda (OpenNext)"],
            ["API", "ECS Fargate — Api.production.eluven.sst:8000"],
            ["DocumentWorker", "ECS Fargate — python -m workers.document_worker"],
            ["WorkflowWorker", "ECS Fargate — python -m workers.workflow_worker"],
            ["Database", "RDS PostgreSQL db.t3.micro, private subnet"],
            ["AI chat", "us.anthropic.claude-sonnet-4-5-20250929-v1:0"],
            ["AI embed", "amazon.titan-embed-text-v1"],
        ],
    )

    add_heading(doc, "8.4 Build sequence — progress (June 2026)", level=2)
    add_table(
        doc,
        ["Weeks", "Focus", "Status", "Notes"],
        [
            ["1–2", "Foundation", "✓ Complete", ""],
            ["3–4", "Document + KB", "◐ Partial", "Workers deployed; .txt verified; PDF path open"],
            ["4–5", "AI core", "✓ Complete", "Verified in prod"],
            ["6–7", "EPR module", "◐ Partial", "Export + UI; needs full task validation + instructions"],
            ["8–9", "SPR module", "◐ Partial", "Assignment UI; batch/grades minimal"],
            ["10–11", "Workflow engine", "◐ Partial", "Workers + UI + seeds; full E2E untested"],
            ["11–12", "Polish", "◐ Partial", "Tests scaffold; RTL + E2E CI remain"],
        ],
    )

    add_heading(doc, "8.5 Production validation checklist", level=2)
    add_table(
        doc,
        ["#", "Flow", "Status"],
        [
            ["1–5", "Login through AI SSE", "✓ Verified"],
            ["6", "KB upload + index", "✓ Verified (.txt)"],
            ["7", "RAG in thread", "○ To test with KB attached to task"],
            ["8–9", "EPR threads + task memory", "○ UI exists; E2E untested"],
            ["10", "Instruction Studio", "○ Partial (3 levels)"],
            ["11", "Hebrew/RTL", "○ To test"],
            ["12", "Workflow + intervention", "○ UI exists; full run untested"],
            ["13", "Word export", "○ Implemented; untested in browser"],
        ],
    )

    add_heading(doc, "11. Decision Log — additions since v1.5", level=2)
    add_table(
        doc,
        ["Date", "Decision", "Rationale"],
        [
            ["Jun 2026", "Dedicated ECS services for document + workflow workers", "Unblocks KB/RAG and workflows; see ADR 001"],
            ["Jun 2026", "Plain-text extraction without Unstructured ML in Docker", "unstructured_inference not in image; fast path for .txt"],
            ["Jun 2026", "Async worker poll loop (single event loop)", "Avoid asyncpg/SQLAlchemy loop conflicts"],
            ["Jun 2026", "Assignment cluster_type + submission child tasks", "SPR hierarchy per domain model"],
        ],
    )

    add_heading(doc, "13. Document Version History", level=1)
    add_table(
        doc,
        ["Version", "Date", "Summary"],
        [
            ["1.6", "June 2026", "Post MVP audit — workers deployed, validation checklist updated, remaining gaps"],
            ["1.5", "June 2026", "Initial production deployment status"],
            ["1.4", "May 2026", "Data model complete; build unblocked"],
        ],
    )

    add_para(doc, "Sections 2, 4–7, 8.1–8.3, 9–10, 12 unchanged in substance from v1.5 / v1.4.", bold=True)

    return doc


def main() -> None:
    doc = build()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
