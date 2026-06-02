"""Generate Word (.docx) exports from Task memory and thread summaries (ADR 002)."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.shared import Pt

from models.memory import TaskMemoryEntry, TaskMemoryEntryType
from models.task import Task
from models.thread import Thread

MODULE_EXTERNAL_PAPER_REVIEW = "external_paper_review"
MODULE_STUDENT_PAPER_REVIEW = "student_paper_review"

_MODULE_EXPORT_COPY: dict[str, dict[str, str]] = {
    MODULE_EXTERNAL_PAPER_REVIEW: {
        "subtitle": "External Paper Review — working export",
        "memory_intro": (
            "Structured findings accumulated during review. "
            "Use alongside thread conversations for final recommendation."
        ),
        "threads_heading": "Review threads",
    },
    MODULE_STUDENT_PAPER_REVIEW: {
        "subtitle": "Student Paper Review — feedback export",
        "memory_intro": (
            "Structured feedback and evaluation notes for this submission. "
            "Use for individual student feedback or grade justification."
        ),
        "threads_heading": "Evaluation threads",
    },
}


def _module_copy(module_type: str) -> dict[str, str]:
    return _MODULE_EXPORT_COPY.get(
        module_type,
        {
            "subtitle": "Task export",
            "memory_intro": "Structured task memory entries.",
            "threads_heading": "Threads",
        },
    )


def _add_memory_section(
    doc: Document,
    heading: str,
    entries: list[TaskMemoryEntry],
) -> None:
    if not entries:
        return
    doc.add_heading(heading, level=2)
    for entry in entries:
        paragraph = doc.add_paragraph(entry.content)
        paragraph.style = "List Bullet"
        if entry.confidence is not None:
            doc.add_paragraph(
                f"Confidence: {entry.confidence:.0%}",
                style="Intense Quote",
            )


def build_task_word_export(
    task: Task,
    memory_entries: list[TaskMemoryEntry],
    threads: list[Thread],
) -> bytes:
    """Build a module-aware Word document summarizing task memory and threads."""
    doc = Document()
    copy = _module_copy(task.module_type)

    title = doc.add_heading(task.title, level=0)
    title.runs[0].font.size = Pt(24)

    doc.add_paragraph(copy["subtitle"])
    doc.add_paragraph(f"Module: {task.module_type}")
    doc.add_paragraph(f"Status: {task.status.value}")

    grouped: dict[TaskMemoryEntryType, list[TaskMemoryEntry]] = {
        TaskMemoryEntryType.FINDING: [],
        TaskMemoryEntryType.ASSUMPTION: [],
        TaskMemoryEntryType.GAP: [],
        TaskMemoryEntryType.REFERENCE: [],
    }
    for entry in memory_entries:
        grouped[entry.entry_type].append(entry)

    doc.add_heading("Task Memory", level=1)
    doc.add_paragraph(copy["memory_intro"])
    _add_memory_section(doc, "Findings", grouped[TaskMemoryEntryType.FINDING])
    _add_memory_section(doc, "Assumptions", grouped[TaskMemoryEntryType.ASSUMPTION])
    _add_memory_section(doc, "Gaps", grouped[TaskMemoryEntryType.GAP])
    _add_memory_section(doc, "References", grouped[TaskMemoryEntryType.REFERENCE])

    if not memory_entries:
        doc.add_paragraph("No memory entries recorded yet.")

    if threads:
        doc.add_heading(copy["threads_heading"], level=1)
        for thread in threads:
            doc.add_heading(thread.title, level=2)
            doc.add_paragraph(f"Thread type: {thread.thread_type}")
            doc.add_paragraph(f"Status: {thread.status.value}")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
