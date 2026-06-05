"""Cross-application document integrity scanning for untrusted uploads."""

from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any

from core.config import settings
from core.logging import get_logger
from schemas.document_integrity import (
    IntegrityFinding,
    IntegrityFindingType,
    IntegrityReport,
    IntegrityStatus,
)

logger = get_logger(__name__)

_INJECTION_PATTERNS: tuple[re.Pattern[str], ...] = (
    re.compile(r"ignore\s+(all\s+)?(previous|prior|above)\s+instructions?", re.I),
    re.compile(r"disregard\s+(all\s+)?(previous|prior|above)", re.I),
    re.compile(r"you\s+are\s+now", re.I),
    re.compile(r"system\s*prompt", re.I),
    re.compile(r"new\s+instructions?\s*:", re.I),
    re.compile(r"do\s+not\s+(note|mention|flag|criticize|grade)", re.I),
    re.compile(r"always\s+(give|assign|recommend|rate)", re.I),
    re.compile(r"grade\s*:\s*[Aa][+]?", re.I),
    re.compile(r"recommend\s+acceptance", re.I),
    re.compile(r"התעלם\s+מ", re.I),
    re.compile(r"הוראות\s+קודמות", re.I),
)

_IMPERATIVE_PATTERN = re.compile(
    r"\b(you must|you should always|never mention|do not flag|always give|ignore all)\b",
    re.I,
)

_UNICODE_SUSPICIOUS = (
    "\u200b",
    "\u200c",
    "\u200d",
    "\ufeff",
    "\u202e",
    "\u202d",
)

_RESERVED_HEADERS = (
    "## task memory",
    "## papers under review",
    "## retrieved knowledge",
    "## task context",
    "## tags",
    "<untrusted_",
    "</untrusted_",
    "```intervention",
)


@dataclass
class DocxRunSignal:
    """Format signal from a DOCX text run."""

    text: str
    hidden: bool = False
    font_size_half_points: int | None = None
    color_rgb: str | None = None


@dataclass
class IntegrityScanInput:
    """Input payload for integrity scanning."""

    text: str
    filename: str
    file_type: str
    docx_run_signals: list[DocxRunSignal] = field(default_factory=list)
    chunk_texts: list[str] | None = None


def content_hash(text: str) -> str:
    """SHA-256 hash of extracted text for tamper detection."""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


class DocumentIntegrityService:
    """Scan extracted document content for integrity risks."""

    def scan(self, payload: IntegrityScanInput) -> IntegrityReport:
        """Run all enabled integrity checks and return a structured report."""
        findings: list[IntegrityFinding] = []
        text = payload.text

        injection_count = self._count_pattern_matches(text, _INJECTION_PATTERNS)
        if injection_count:
            findings.append(
                IntegrityFinding(
                    finding_type=IntegrityFindingType.PROMPT_INJECTION_PATTERN,
                    count=injection_count,
                    detail="Instruction-override language detected in document text",
                ),
            )

        hidden_count = self._scan_hidden_formatting(payload)
        if hidden_count:
            findings.append(
                IntegrityFinding(
                    finding_type=IntegrityFindingType.HIDDEN_FORMATTING,
                    count=hidden_count,
                    detail="Hidden or abnormally formatted text runs detected",
                ),
            )

        unicode_count = self._scan_unicode_anomalies(text)
        if unicode_count:
            findings.append(
                IntegrityFinding(
                    finding_type=IntegrityFindingType.UNICODE_ANOMALY,
                    count=unicode_count,
                    detail="Suspicious Unicode control or zero-width characters detected",
                ),
            )

        mimicry_count = self._scan_structural_mimicry(text)
        if mimicry_count:
            findings.append(
                IntegrityFinding(
                    finding_type=IntegrityFindingType.STRUCTURAL_MIMICRY,
                    count=mimicry_count,
                    detail="Document text mimics platform context structure",
                ),
            )

        filename_count = self._scan_filename(payload.filename)
        if filename_count:
            findings.append(
                IntegrityFinding(
                    finding_type=IntegrityFindingType.FILENAME_SUSPICIOUS,
                    count=filename_count,
                    detail="Filename contains instruction-like language",
                ),
            )

        stuffing_count = self._scan_context_stuffing(text, payload.docx_run_signals)
        if stuffing_count:
            findings.append(
                IntegrityFinding(
                    finding_type=IntegrityFindingType.CONTEXT_STUFFING,
                    count=stuffing_count,
                    detail="Abnormally large or hidden text volume detected",
                ),
            )

        density_count = self._scan_instruction_density(text)
        if density_count:
            findings.append(
                IntegrityFinding(
                    finding_type=IntegrityFindingType.INSTRUCTION_DENSITY,
                    count=density_count,
                    detail="High density of imperative instruction-like sentences",
                ),
            )

        if payload.chunk_texts:
            rag_bait_count = self._scan_rag_bait(payload.chunk_texts)
            if rag_bait_count:
                findings.append(
                    IntegrityFinding(
                        finding_type=IntegrityFindingType.RAG_BAIT,
                        count=rag_bait_count,
                        detail="Repeated directive language across document chunks",
                    ),
                )

        status = self._aggregate_status(findings)
        token_count = len(text.split()) if text else 0

        report = IntegrityReport(
            status=status,
            findings=findings,
            content_hash=content_hash(text),
            scanned_at=datetime.now(UTC),
            extract_token_count=token_count,
        )

        logger.info(
            "document_integrity_scan_complete",
            filename=payload.filename,
            file_type=payload.file_type,
            status=status.value,
            finding_count=len(findings),
        )
        return report

    @staticmethod
    def _count_pattern_matches(text: str, patterns: tuple[re.Pattern[str], ...]) -> int:
        total = 0
        for pattern in patterns:
            total += len(pattern.findall(text))
        return total

    def _scan_hidden_formatting(self, payload: IntegrityScanInput) -> int:
        if payload.file_type != "docx" or not payload.docx_run_signals:
            return 0

        count = 0
        min_font = settings.integrity_min_font_half_points
        for run in payload.docx_run_signals:
            if run.hidden:
                count += 1
                continue
            if run.font_size_half_points is not None and run.font_size_half_points < min_font:
                count += 1
            if run.color_rgb == "FFFFFF" and run.text.strip():
                count += 1
        return count

    @staticmethod
    def _scan_unicode_anomalies(text: str) -> int:
        count = 0
        for char in _UNICODE_SUSPICIOUS:
            count += text.count(char)
        for char in text:
            if unicodedata.category(char) in {"Cf", "Co", "Cs"} and char not in "\n\r\t":
                count += 1
        return count

    @staticmethod
    def _scan_structural_mimicry(text: str) -> int:
        lowered = text.lower()
        count = 0
        for header in _RESERVED_HEADERS:
            count += lowered.count(header)
        return count

    def _scan_filename(self, filename: str) -> int:
        name = filename.rsplit(".", 1)[0] if "." in filename else filename
        normalized = name.replace("-", " ").replace("_", " ")
        return self._count_pattern_matches(normalized, _INJECTION_PATTERNS)

    def _scan_context_stuffing(
        self,
        text: str,
        docx_run_signals: list[DocxRunSignal],
    ) -> int:
        if len(text) < settings.integrity_stuffing_min_chars:
            return 0

        hidden_chars = sum(len(run.text) for run in docx_run_signals if run.hidden)
        if hidden_chars >= settings.integrity_stuffing_hidden_chars:
            return 1

        if len(text) >= settings.integrity_stuffing_max_chars:
            return 1

        return 0

    @staticmethod
    def _scan_instruction_density(text: str) -> int:
        sentences = re.split(r"[.!?]\s+", text)
        if len(sentences) < 5:
            return 0
        imperative = sum(1 for sentence in sentences if _IMPERATIVE_PATTERN.search(sentence))
        ratio = imperative / len(sentences)
        if ratio >= settings.integrity_imperative_density_threshold:
            return imperative
        return 0

    @staticmethod
    def _scan_rag_bait(chunk_texts: list[str]) -> int:
        if len(chunk_texts) < 2:
            return 0
        directive_chunks = 0
        for chunk in chunk_texts:
            if _IMPERATIVE_PATTERN.search(chunk) or any(
                pattern.search(chunk) for pattern in _INJECTION_PATTERNS[:4]
            ):
                directive_chunks += 1
        if directive_chunks >= 2:
            return directive_chunks
        return 0

    @staticmethod
    def _aggregate_status(findings: list[IntegrityFinding]) -> IntegrityStatus:
        if not findings:
            return IntegrityStatus.CLEAN

        critical_types = {
            IntegrityFindingType.PROMPT_INJECTION_PATTERN,
            IntegrityFindingType.STRUCTURAL_MIMICRY,
            IntegrityFindingType.HIDDEN_FORMATTING,
        }
        if any(finding.finding_type in critical_types for finding in findings):
            return IntegrityStatus.REVIEW_REQUIRED

        total_count = sum(finding.count for finding in findings)
        if total_count >= settings.integrity_warning_threshold:
            return IntegrityStatus.REVIEW_REQUIRED
        return IntegrityStatus.WARNING


def extract_docx_run_signals(file_bytes: bytes) -> list[DocxRunSignal]:
    """Extract per-run formatting signals from DOCX for integrity scanning."""
    import io

    from docx import Document

    signals: list[DocxRunSignal] = []
    document = Document(io.BytesIO(file_bytes))

    def inspect_runs(paragraph: Any) -> None:
        for run in paragraph.runs:
            text = run.text
            if not text:
                continue
            hidden = bool(run.font.hidden)
            size = run.font.size.pt * 2 if run.font.size is not None else None
            color_rgb: str | None = None
            if run.font.color is not None and run.font.color.rgb is not None:
                color_rgb = str(run.font.color.rgb).upper()
            signals.append(
                DocxRunSignal(
                    text=text,
                    hidden=hidden,
                    font_size_half_points=int(size) if size is not None else None,
                    color_rgb=color_rgb,
                ),
            )

    for paragraph in document.paragraphs:
        inspect_runs(paragraph)

    for section in document.sections:
        for header_paragraph in section.header.paragraphs:
            inspect_runs(header_paragraph)
        for footer_paragraph in section.footer.paragraphs:
            inspect_runs(footer_paragraph)

    return signals
