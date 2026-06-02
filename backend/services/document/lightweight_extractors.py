"""MVP document extractors without Unstructured ML dependencies (ADR 002)."""

from __future__ import annotations

import csv
import io
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from core.logging import get_logger

logger = get_logger(__name__)


def extract_txt_or_tex(file_bytes: bytes, filename: str, file_type: str) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    """Decode plain text or LaTeX source."""
    text = file_bytes.decode("utf-8", errors="replace").strip()
    metadata = _base_metadata(filename, file_type, len(file_bytes), element_count=1 if text else 0)
    return text, [], metadata


def extract_pdf(file_bytes: bytes, filename: str) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    """Extract text from PDF via pdfminer (text layer only — no OCR)."""
    from pdfminer.high_level import extract_text

    text = extract_text(io.BytesIO(file_bytes)).strip()
    metadata = _base_metadata(filename, "pdf", len(file_bytes), element_count=1 if text else 0)
    metadata["extractor"] = "pdfminer.six"
    metadata["ocr"] = False
    return text, [], metadata


def extract_docx(file_bytes: bytes, filename: str) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    """Extract paragraphs and table cell text from DOCX."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(file_bytes))
    text_parts: list[str] = []
    tables: list[dict[str, Any]] = []

    for block in _iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            snippet = block.text.strip()
            if snippet:
                text_parts.append(snippet)
        elif isinstance(block, Table):
            rows: list[list[str]] = []
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows.append(cells)
            if rows:
                table_text = "\n".join(" | ".join(row) for row in rows)
                tables.append({"text": table_text, "metadata": {"row_count": len(rows)}})
                text_parts.append(table_text)

    text = "\n\n".join(text_parts)
    metadata = _base_metadata(filename, "docx", len(file_bytes), element_count=len(text_parts))
    metadata["extractor"] = "python-docx"
    return text, tables, metadata


def extract_xlsx(file_bytes: bytes, filename: str) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    """Extract cell values from Excel as text blocks per sheet."""
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    text_parts: list[str] = []
    tables: list[dict[str, Any]] = []

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell).strip() if cell is not None else "" for cell in row]
            if any(cells):
                rows.append(cells)
        if rows:
            sheet_text = f"Sheet: {sheet_name}\n" + "\n".join(" | ".join(r) for r in rows)
            text_parts.append(sheet_text)
            tables.append({"text": sheet_text, "metadata": {"sheet": sheet_name, "row_count": len(rows)}})

    workbook.close()
    text = "\n\n".join(text_parts)
    metadata = _base_metadata(filename, "xlsx", len(file_bytes), element_count=len(text_parts))
    metadata["extractor"] = "openpyxl"
    return text, tables, metadata


def extract_csv(file_bytes: bytes, filename: str) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    """Parse CSV as delimited text."""
    decoded = file_bytes.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(decoded))
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    text = "\n".join(" | ".join(row) for row in rows)
    tables = [{"text": text, "metadata": {"row_count": len(rows)}}] if rows else []
    metadata = _base_metadata(filename, "csv", len(file_bytes), element_count=len(rows))
    metadata["extractor"] = "csv"
    return text, tables, metadata


def extract_doc(file_bytes: bytes, filename: str) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    """Convert legacy .doc to text via LibreOffice headless (in runtime image)."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        input_path = Path(tmp_dir) / filename
        input_path.write_bytes(file_bytes)
        result = subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "txt:Text",
                "--outdir",
                tmp_dir,
                str(input_path),
            ],
            capture_output=True,
            text=True,
            check=False,
            timeout=120,
        )
        if result.returncode != 0:
            logger.warning(
                "doc_conversion_failed",
                filename=filename,
                stderr=result.stderr[:500],
            )
            raise ValueError(f"Could not convert legacy Word document: {filename}")

        output_path = input_path.with_suffix(".txt")
        if not output_path.exists():
            raise ValueError(f"LibreOffice produced no output for: {filename}")

        text = output_path.read_text(encoding="utf-8", errors="replace").strip()

    metadata = _base_metadata(filename, "doc", len(file_bytes), element_count=1 if text else 0)
    metadata["extractor"] = "libreoffice"
    return text, [], metadata


def _base_metadata(filename: str, file_type: str, size_bytes: int, *, element_count: int) -> dict[str, Any]:
    return {
        "filename": filename,
        "file_type": file_type,
        "size_bytes": size_bytes,
        "element_count": element_count,
        "element_types": ["Text"],
    }


def _iter_docx_blocks(document: Any):
    """Yield paragraphs and tables in document order."""
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parent = document.element.body
    for child in parent.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)
